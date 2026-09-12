import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.agent.workflows import InspectionToApprovalWorkflow

BASE = Path(__file__).resolve().parent.parent
wf = InspectionToApprovalWorkflow()

print("#" * 74)
print("# SCRIPTED WORKFLOW: inspection report -> SOP lookup -> approval note")
print("#" * 74)


def show(st):
    print(f"  [{st.n}] {'ok ' if st.ok else 'ERR'} {st.name:<28} "
          f"{st.elapsed_ms:>6} ms   {st.detail[:110]}")


run = wf.run("data/samples/inspection_report.png",
             "workflow_approval.docx", on_stage=show)

print(f"\n  total     : {run.total_ms} ms")
print(f"  artifacts : {run.artifacts}")
print(f"  failed    : {run.failed or 'no'}")
print(f"\n  {run.summary}")

target = BASE / "outputs" / "workflow_approval.docx"
print("\n--- DOCUMENT FACT CHECK ---")
if not target.exists():
    print(f"  MISSING FILE: {target}")
    sys.exit(1)

doc = Document(target)
body = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
for t in doc.tables:
    for row in t.rows:
        body += "\n" + " | ".join(c.text for c in row.cells)

FACTS = {
    "P-101A": "equipment tag",
    "8.2": "measured casing thickness",
    "9.0": "design minimum",
    "82": "bearing temperature",
    "75": "alarm setpoint",
    "INSP-2026-0417": "source report number",
    "30": "shutdown window (from SOP)",
    "Mechanical Maintenance": "approving authority (from SOP)",
    "SOP-MECH-014": "SOP citation",
}
missing = []
for token, meaning in FACTS.items():
    ok = token.lower() in body.lower()
    print(f"  {'FOUND  ' if ok else 'MISSING'}  {token:<24} {meaning}")
    if not ok:
        missing.append(token)

print(f"\n  {len(FACTS)-len(missing)}/{len(FACTS)} facts present")
print("\n--- DOCUMENT BODY ---")
print(body[:1800])
