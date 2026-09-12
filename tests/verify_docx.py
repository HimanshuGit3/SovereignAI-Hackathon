import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

OUT = Path(__file__).resolve().parent.parent / "outputs"
target = OUT / "approval_P101A.docx"

if not target.exists():
    print(f"MISSING: {target}")
    sys.exit(1)

doc = Document(target)
full = []
for p in doc.paragraphs:
    if p.text.strip():
        full.append(p.text)
for t in doc.tables:
    for row in t.rows:
        full.append(" | ".join(c.text for c in row.cells))

body = "\n".join(full)
print("=" * 70)
print(body)
print("=" * 70)

print("\n--- FACT CHECK against the source report ---")
FACTS = {
    "P-101A": "equipment tag",
    "8.2": "measured casing thickness (mm)",
    "9.0": "design minimum (mm)",
    "82": "bearing temperature (deg C)",
    "75": "alarm setpoint (deg C)",
    "INSP-2026-0417": "source report number",
    "30 days": "shutdown window",
}
missing = []
for token, meaning in FACTS.items():
    present = token in body
    print(f"  {'FOUND  ' if present else 'MISSING'}  {token:<16} {meaning}")
    if not present:
        missing.append(token)

print(f"\n  {len(FACTS)-len(missing)}/{len(FACTS)} source facts preserved")
if missing:
    print(f"  absent: {missing}")
