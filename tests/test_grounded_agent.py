import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.agent.loop import Agent

agent = Agent()

TASK = (
    "Read the inspection report at data/samples/inspection_report.png. "
    "Check our SOPs for who must approve a shutdown at this criticality "
    "and how many days we have. Then draft a formal approval note as a "
    "Word document called grounded_approval.docx citing the SOP."
)

print("#" * 74)
print("# GROUNDED FLAGSHIP: scan -> SOP lookup -> document")
print("#" * 74)
print(f"\nTASK: {TASK}\n")


def show(step):
    if step.kind == "tool":
        flag = "ok " if step.ok else "ERR"
        print(f"  [{step.n}] {flag} {step.tool}  {step.elapsed_ms} ms")
        print(f"       args: {str(step.arguments)[:150]}")
        print(f"       obs:  {step.observation[:220].replace(chr(10),' ')}")
    else:
        print(f"  [{step.n}] FINAL  {step.elapsed_ms} ms")


run = agent.run(TASK, on_step=show)

print(f"\n  tools offered : {run.tools_offered}")
print(f"  tools used    : {[s.tool for s in run.steps if s.kind=='tool']}")
print(f"  total         : {run.total_ms} ms")
print(f"  artifacts     : {run.artifacts}")
print(f"\n  ANSWER:\n    {run.answer[:700].replace(chr(10), chr(10)+'    ')}")

print("\n--- verifying the generated document ---")
from docx import Document  # noqa: E402
target = Path(__file__).resolve().parent.parent / "outputs" / "grounded_approval.docx"
if target.exists():
    body = "\n".join(p.text for p in Document(target).paragraphs if p.text.strip())
    for t in Document(target).tables:
        for row in t.rows:
            body += "\n" + " | ".join(c.text for c in row.cells)
    FACTS = {"P-101A": "equipment tag", "8.2": "measured thickness",
             "9.0": "design minimum", "82": "bearing temp",
             "INSP-2026-0417": "source report", "30": "shutdown window days",
             "Mechanical Maintenance": "approving authority from SOP"}
    for token, meaning in FACTS.items():
        print(f"  {'FOUND  ' if token in body else 'MISSING'}  {token:<22} {meaning}")
else:
    print(f"  MISSING: {target}")

print("\n=== GROUNDED AGENT TEST COMPLETE ===")
