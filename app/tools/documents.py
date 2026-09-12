"""Generate real .docx deliverables, not chat replies."""
import json
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.config import BASE_DIR
from app.tools.base import ToolResult, ToolSpec, safe_path


def _coerce_list(value) -> list[str]:
    """Models pass lists as JSON strings, newline text, or real lists."""
    if value is None:
        return []
    if isinstance(value, list):
        return [str(v).strip() for v in value if str(v).strip()]
    if isinstance(value, str):
        v = value.strip()
        if v.startswith("["):
            try:
                parsed = json.loads(v)
                if isinstance(parsed, list):
                    return [str(x).strip() for x in parsed if str(x).strip()]
            except json.JSONDecodeError:
                pass
        return [ln.strip(" -*\t") for ln in v.splitlines() if ln.strip(" -*\t")]
    return [str(value)]


def create_word_document(
    filename: str,
    title: str,
    subject: str = "",
    reference: str = "",
    background: str = "",
    findings=None,
    recommendation: str = "",
    approval_sought: str = "",
    prepared_by: str = "SovereignAI Workbench",
) -> ToolResult:
    name = filename if filename.lower().endswith(".docx") else f"{filename}.docx"
    out = safe_path(f"outputs/{name}")
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(15)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mrun = meta.add_run(
        f"Generated on-premise  |  {datetime.now().strftime('%d %B %Y, %H:%M')}")
    mrun.font.size = Pt(8)
    mrun.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    if reference or subject:
        doc.add_paragraph()
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for label, value in (("Reference", reference), ("Subject", subject)):
            if value:
                cells = table.add_row().cells
                cells[0].paragraphs[0].add_run(label).bold = True
                cells[1].text = value

    def section(heading: str, body: str):
        if not body:
            return
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.add_run(heading.upper()).bold = True
        doc.add_paragraph(body)

    section("Background", background)

    items = _coerce_list(findings)
    if items:
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.add_run("FINDINGS").bold = True
        for item in items:
            doc.add_paragraph(item, style="List Number")

    section("Recommendation", recommendation)
    section("Approval Sought", approval_sought)

    doc.add_paragraph()
    doc.add_paragraph()
    sig = doc.add_paragraph()
    sig.add_run("Prepared by: ").bold = True
    sig.add_run(prepared_by)
    doc.add_paragraph("_" * 34)
    doc.add_paragraph("Approving Authority          Date")

    foot = doc.add_paragraph()
    frun = foot.add_run(
        "Produced by an air-gapped local AI system. No data left the premises.")
    frun.font.size = Pt(7)
    frun.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.save(out)
    size_kb = round(out.stat().st_size / 1024, 1)
    return ToolResult(
        True,
        f"Created Word document '{out.name}' ({size_kb} KB) with "
        f"{len(items)} numbered finding(s).",
        "create_word_document",
        artifacts=[str(out.relative_to(BASE_DIR))],
    )


SPECS = [
    ToolSpec(
        name="create_word_document",
        description=(
            "Create a formal .docx approval note or office note in outputs/. "
            "Use this whenever the user asks for a Word file, approval note, "
            "office note or formal document deliverable."
        ),
        parameters={
            "type": "object",
            "properties": {
                "filename": {"type": "string", "description": "e.g. approval_note.docx"},
                "title": {"type": "string", "description": "Document title"},
                "subject": {"type": "string", "description": "One-line subject"},
                "reference": {"type": "string", "description": "Source reference number"},
                "background": {"type": "string", "description": "Context paragraph"},
                "findings": {"type": "array", "items": {"type": "string"},
                             "description": "Numbered findings"},
                "recommendation": {"type": "string", "description": "Recommended action"},
                "approval_sought": {"type": "string",
                                    "description": "What approval is requested"},
            },
            "required": ["filename", "title"],
        },
        fn=create_word_document,
    ),
]
